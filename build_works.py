"""Builds data.js for the Ongoing Works dashboard from the tracker workbook.
Usage:  python build_works.py Departmental_Ongoing_Works_Professional.xlsx
Needs:  pip install openpyxl
Every sheet except 'Executive Summary' is read: header on row 3, tasks from row 4
(Sr. #, area, scope, status, owner, remarks)."""
import sys, re, json
from openpyxl import load_workbook
src = sys.argv[1] if len(sys.argv) > 1 else 'Departmental_Ongoing_Works_Professional.xlsx'
wb = load_workbook(src, data_only=True)
tasks, asof = [], ''
for ws in wb.worksheets:
    if ws.title.strip().lower() == 'executive summary': continue
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) > 1 and isinstance(rows[1][0], str):
        m = re.search(r'Status as of:\s*(.+)$', rows[1][0])
        if m: asof = m.group(1).strip()
    for r in rows[3:]:
        if isinstance(r[0], (int, float)) and r[3]:
            tasks.append({'dept': ws.title, 'sr': int(r[0]), 'area': r[1], 'scope': r[2],
                          'status': str(r[3]).strip(), 'owner': r[4], 'remarks': r[5]})
open('data.js', 'w', encoding='utf-8').write('window.WORKS_DATA=' + json.dumps({'asof': asof, 'tasks': tasks}, ensure_ascii=False, separators=(',', ':')) + ';\n')
print(f'data.js written: {len(tasks)} tasks, status as of {asof}')
